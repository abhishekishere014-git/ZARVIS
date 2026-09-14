"""Native Microsoft PowerPoint (.pptx) generation tool with structural verification."""

from typing import Any, Dict, List, Optional
import pptx
from pptx.util import Inches, Pt
from jarvis.tools.builtins.office.base import (
    DEFAULT_LIMITS,
    verify_file_size_within_limit,
)
from jarvis.tools.decorator import tool
from jarvis.tools.errors import (
    ArtifactVerificationError,
    ResourceLimitExceededError,
)
from jarvis.tools.models import (
    RiskLevel,
    ToolExecutionContext,
    ToolPermission,
)


@tool(
    name="create_pptx",
    tool_id="office.create_pptx",
    description="Generate a formatted Microsoft PowerPoint (.pptx) presentation with title slide, content slides, and bullet points.",
    category="office",
    version="1.0.0",
    permissions={ToolPermission.WRITE},
    risk_level=RiskLevel.LOW,
    timeout_seconds=30.0,
)
async def create_pptx(
    filename: str,
    title: str,
    slides: List[Dict[str, Any]],
    subtitle: Optional[str] = None,
    context: Optional[ToolExecutionContext] = None,
) -> Dict[str, Any]:
    """Generates a PowerPoint presentation (.pptx) and verifies its internal structure.

    :param filename: Output filename, e.g. 'briefing.pptx'
    :param title: Main title for presentation
    :param slides: List of slide dicts (keys: 'title', 'bullets', 'content')
    :param subtitle: Optional subtitle for title slide
    """
    if not context or not context.sandbox:
        raise ValueError("Tool execution context with sandbox is required.")

    # 1. Resolve safe destination inside workspace
    clean_name = context.sandbox.sanitize_filename(filename)
    if not clean_name.lower().endswith(".pptx"):
        clean_name += ".pptx"

    safe_path = context.sandbox.resolve_safe_path(clean_name, allowed_extensions={"pptx"})

    # Check slide limits
    total_slides = 1 + len(slides)  # 1 title slide + content slides
    if total_slides > DEFAULT_LIMITS.max_slides:
        raise ResourceLimitExceededError(
            f"Slide count ({total_slides}) exceeds maximum allowed of {DEFAULT_LIMITS.max_slides}"
        )

    # 2. Build Presentation
    prs = pptx.Presentation()

    # Title Slide
    title_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_layout)
    title_slide.shapes.title.text = title
    if subtitle and len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = subtitle

    # Content Slides
    bullet_layout = prs.slide_layouts[1]
    for slide_spec in slides:
        slide_title = str(slide_spec.get("title", ""))
        bullets = slide_spec.get("bullets", [])
        content = slide_spec.get("content")

        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_title

        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()

        if content:
            p = tf.paragraphs[0]
            p.text = str(content)

        if bullets and isinstance(bullets, list):
            for b_idx, bullet_text in enumerate(bullets):
                p = tf.add_paragraph() if (content or b_idx > 0) else tf.paragraphs[0]
                p.text = str(bullet_text)
                p.level = 0

    # 3. Save presentation
    prs.save(str(safe_path))

    # 4. Verify artifact size and structural integrity
    size_bytes = verify_file_size_within_limit(safe_path)
    try:
        reopened = pptx.Presentation(str(safe_path))
        actual_slide_count = len(reopened.slides)
        if actual_slide_count != total_slides:
            raise ArtifactVerificationError(
                f"Expected {total_slides} slides in presentation, but verified {actual_slide_count}"
            )
    except Exception as exc:
        raise ArtifactVerificationError(f"PPTX structural verification failed: {exc}") from exc

    return {
        "status": "success",
        "filename": safe_path.name,
        "path": str(safe_path),
        "size_bytes": size_bytes,
        "slides_count": actual_slide_count,
        "artifacts": [str(safe_path)],
    }
