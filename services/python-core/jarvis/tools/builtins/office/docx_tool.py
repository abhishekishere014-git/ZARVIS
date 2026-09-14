"""Native Microsoft Word (.docx) generation tool with structural verification."""

from typing import Any, Dict, List, Optional
import docx
from jarvis.tools.builtins.office.base import (
    DEFAULT_LIMITS,
    verify_file_size_within_limit,
)
from jarvis.tools.decorator import tool
from jarvis.tools.errors import (
    ArtifactVerificationError,
    ResourceLimitExceededError,
)
from jarvis.tools.models import (
    RiskLevel,
    ToolExecutionContext,
    ToolPermission,
)


@tool(
    name="create_docx",
    tool_id="office.create_docx",
    description="Generate a professional Microsoft Word (.docx) document with titles, sections, paragraphs, bullet lists, and tables.",
    category="office",
    version="1.0.0",
    permissions={ToolPermission.WRITE},
    risk_level=RiskLevel.LOW,
    timeout_seconds=30.0,
)
async def create_docx(
    filename: str,
    title: str,
    sections: List[Dict[str, Any]],
    subtitle: Optional[str] = None,
    context: Optional[ToolExecutionContext] = None,
) -> Dict[str, Any]:
    """Generates a Microsoft Word document (.docx) and verifies its internal structure.

    :param filename: Output filename, e.g. 'report.docx'
    :param title: Main document title
    :param sections: List of section dicts (keys: 'heading', 'level', 'paragraph', 'bullets', 'table')
    :param subtitle: Optional subtitle below main title
    """
    if not context or not context.sandbox:
        raise ValueError("Tool execution context with sandbox is required.")

    # 1. Resolve safe destination inside workspace
    clean_name = context.sandbox.sanitize_filename(filename)
    if not clean_name.lower().endswith(".docx"):
        clean_name += ".docx"

    safe_path = context.sandbox.resolve_safe_path(clean_name, allowed_extensions={"docx"})

    # 2. Build Document
    doc = docx.Document()
    doc.add_heading(title, level=0)

    if subtitle:
        p_sub = doc.add_paragraph()
        run = p_sub.add_run(subtitle)
        run.italic = True

    total_rows = 0
    for sec in sections:
        heading = sec.get("heading")
        if heading:
            doc.add_heading(str(heading), level=int(sec.get("level", 1)))

        para = sec.get("paragraph")
        if para:
            doc.add_paragraph(str(para))

        bullets = sec.get("bullets")
        if bullets and isinstance(bullets, list):
            for bullet in bullets:
                doc.add_paragraph(str(bullet), style="List Bullet")

        table_data = sec.get("table")
        if table_data and isinstance(table_data, list) and len(table_data) > 0:
            total_rows += len(table_data)
            if total_rows > DEFAULT_LIMITS.max_table_rows:
                raise ResourceLimitExceededError(
                    f"Total table rows ({total_rows}) exceeds limit of {DEFAULT_LIMITS.max_table_rows}"
                )

            headers = table_data[0]
            num_cols = len(headers)
            if num_cols > DEFAULT_LIMITS.max_table_cols:
                raise ResourceLimitExceededError(
                    f"Table columns ({num_cols}) exceeds limit of {DEFAULT_LIMITS.max_table_cols}"
                )

            table = doc.add_table(rows=1, cols=num_cols)
            table.style = "Table Grid"
            hdr_cells = table.rows[0].cells
            for idx, hdr in enumerate(headers):
                hdr_cells[idx].text = str(hdr)

            for row_data in table_data[1:]:
                row_cells = table.add_row().cells
                for idx, cell_value in enumerate(row_data[:num_cols]):
                    row_cells[idx].text = str(cell_value)

    # 3. Save to file
    doc.save(str(safe_path))

    # 4. Verify artifact size and structural integrity
    size_bytes = verify_file_size_within_limit(safe_path)
    try:
        reopened = docx.Document(str(safe_path))
        num_paras = len(reopened.paragraphs)
        num_tables = len(reopened.tables)
        if num_paras == 0 and num_tables == 0:
            raise ArtifactVerificationError("Reopened DOCX document contains no paragraphs or tables.")
    except Exception as exc:
        raise ArtifactVerificationError(f"DOCX structural verification failed: {exc}") from exc

    return {
        "status": "success",
        "filename": safe_path.name,
        "path": str(safe_path),
        "size_bytes": size_bytes,
        "paragraphs_count": num_paras,
        "tables_count": num_tables,
        "artifacts": [str(safe_path)],
    }
