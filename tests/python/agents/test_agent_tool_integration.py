"""Tests for Agent Runtime integration with Phase 04 ToolExecutor and Office generation."""

from pathlib import Path
import docx
import openpyxl
import pytest
from jarvis.agents.base import AgentContext
from jarvis.agents.builtins.coding import CodingAgent
from jarvis.agents.models import AgentTask, TaskState
from jarvis.agents.verifier import AgentVerifier
from jarvis.tools.factory import build_tool_system


@pytest.mark.asyncio
async def test_coding_agent_generates_docx_via_tool_executor(tmp_path: Path) -> None:
    tool_system = build_tool_system(workspace_dir=tmp_path)
    coder = CodingAgent()
    verifier = AgentVerifier(sandbox=tool_system.sandbox)

    task = AgentTask(
        id="task_docx_gen",
        title="Quarterly Operations Review",
        description="Comprehensive analysis of operational throughput and latency",
        assigned_agent=coder.id,
        required_tools=["office.create_docx"],
    )

    ctx = AgentContext(
        run_id="run_docx_test",
        task=task,
        agent_definition=coder.definition,
        tool_executor=tool_system.executor,
    )

    res = await coder.execute(ctx)
    assert res.status == TaskState.COMPLETED
    assert len(res.artifacts) == 1

    # Verify via verifier
    ver = verifier.verify_step(task, res)
    assert ver.state.value == "passed"
    assert len(ver.verified_artifacts) == 1

    # Verify physical file structure
    doc_path = Path(res.artifacts[0])
    assert doc_path.exists()
    doc = docx.Document(str(doc_path))
    assert len(doc.paragraphs) > 0


@pytest.mark.asyncio
async def test_coding_agent_generates_xlsx_via_tool_executor(tmp_path: Path) -> None:
    tool_system = build_tool_system(workspace_dir=tmp_path)
    coder = CodingAgent()
    verifier = AgentVerifier(sandbox=tool_system.sandbox)

    task = AgentTask(
        id="task_xlsx_gen",
        title="Q3 Budget Matrix",
        description="Spreadsheet of allocation by business unit",
        assigned_agent=coder.id,
        required_tools=["office.create_xlsx"],
    )

    ctx = AgentContext(
        run_id="run_xlsx_test",
        task=task,
        agent_definition=coder.definition,
        tool_executor=tool_system.executor,
    )

    res = await coder.execute(ctx)
    assert res.status == TaskState.COMPLETED
    assert len(res.artifacts) == 1

    ver = verifier.verify_step(task, res)
    assert ver.state.value == "passed"

    xlsx_path = Path(res.artifacts[0])
    assert xlsx_path.exists()
    wb = openpyxl.load_workbook(str(xlsx_path), read_only=True)
    assert len(wb.sheetnames) >= 1
    wb.close()
