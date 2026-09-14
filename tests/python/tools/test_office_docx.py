"""Tests for native DOCX generation and structural verification."""

from pathlib import Path
import docx
import pytest
from jarvis.tools.builtins.office.docx_tool import create_docx
from jarvis.tools.errors import ArtifactVerificationError, ResourceLimitExceededError
from jarvis.tools.models import ToolExecutionContext
from jarvis.tools.sandbox import FileSystemSandbox


@pytest.mark.asyncio
async def test_create_docx_full_lifecycle(tmp_path: Path) -> None:
    sandbox = FileSystemSandbox(workspace_root=tmp_path)
    context = ToolExecutionContext(
        execution_id="test_docx_1",
        tool_id="office.create_docx",
        workspace_dir=tmp_path,
        sandbox=sandbox,
    )

    sections = [
        {"heading": "Executive Summary", "level": 1, "paragraph": "This is an automated intelligence briefing."},
        {"heading": "Key Points", "level": 2, "bullets": ["First finding", "Second finding", "Third finding"]},
        {
            "heading": "Metrics Table",
            "level": 2,
            "table": [
                ["Component", "Latency", "Status"],
                ["AI Provider", "0.26ms", "Optimal"],
                ["Tool Engine", "0.45ms", "Optimal"],
            ],
        },
    ]

    result = await create_docx(
        filename="briefing.docx",
        title="JARVIS Intelligence Briefing",
        subtitle="Confidential System Assessment",
        sections=sections,
        context=context,
    )

    assert result["status"] == "success"
    assert result["filename"] == "briefing.docx"
    assert result["size_bytes"] > 0
    assert result["paragraphs_count"] >= 5
    assert result["tables_count"] == 1

    # Verify physical file existence and valid docx structure
    file_path = Path(result["path"])
    assert file_path.exists()

    opened = docx.Document(str(file_path))
    assert len(opened.paragraphs) >= 5
    assert len(opened.tables) == 1
    assert opened.tables[0].rows[0].cells[0].text == "Component"


@pytest.mark.asyncio
async def test_create_docx_enforces_limits(tmp_path: Path) -> None:
    sandbox = FileSystemSandbox(workspace_root=tmp_path)
    context = ToolExecutionContext(
        execution_id="test_docx_2",
        tool_id="office.create_docx",
        workspace_dir=tmp_path,
        sandbox=sandbox,
    )

    # Exceed table row limit
    huge_table = [["Col1", "Col2"]] + [["val1", "val2"] for _ in range(10_005)]
    sections = [{"heading": "Huge Table", "table": huge_table}]

    with pytest.raises(ResourceLimitExceededError):
        await create_docx(
            filename="overflow.docx",
            title="Overflow",
            sections=sections,
            context=context,
        )
