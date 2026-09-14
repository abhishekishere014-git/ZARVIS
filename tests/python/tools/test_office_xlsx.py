"""Tests for native XLSX generation and workbook verification."""

from pathlib import Path
import openpyxl
import pytest
from jarvis.tools.builtins.office.xlsx_tool import create_xlsx
from jarvis.tools.errors import ResourceLimitExceededError
from jarvis.tools.models import ToolExecutionContext
from jarvis.tools.sandbox import FileSystemSandbox


@pytest.mark.asyncio
async def test_create_xlsx_full_lifecycle(tmp_path: Path) -> None:
    sandbox = FileSystemSandbox(workspace_root=tmp_path)
    context = ToolExecutionContext(
        execution_id="test_xlsx_1",
        tool_id="office.create_xlsx",
        workspace_dir=tmp_path,
        sandbox=sandbox,
    )

    sheets = [
        {
            "name": "Revenue",
            "headers": ["Quarter", "Target", "Actual", "Growth %"],
            "rows": [
                ["Q1 2026", 100000, 115000, 15.0],
                ["Q2 2026", 120000, 128000, 6.67],
                ["Q3 2026", 140000, 152000, 8.57],
            ],
        },
        {
            "name": "Expenses",
            "headers": ["Category", "Budget", "Spent"],
            "rows": [
                ["Compute", 20000, 18500],
                ["R&D", 50000, 49200],
            ],
        },
    ]

    result = await create_xlsx(
        filename="financial_model.xlsx",
        sheets=sheets,
        context=context,
    )

    assert result["status"] == "success"
    assert result["filename"] == "financial_model.xlsx"
    assert result["sheets_count"] == 2
    assert result["sheet_names"] == ["Revenue", "Expenses"]
    assert result["total_rows"] == 5

    # Verify physical file and workbook structure
    file_path = Path(result["path"])
    assert file_path.exists()

    wb = openpyxl.load_workbook(str(file_path))
    assert wb.sheetnames == ["Revenue", "Expenses"]
    rev_sheet = wb["Revenue"]
    assert rev_sheet.cell(row=1, column=1).value == "Quarter"
    assert rev_sheet.cell(row=2, column=2).value == 100000
    wb.close()


@pytest.mark.asyncio
async def test_create_xlsx_enforces_limits(tmp_path: Path) -> None:
    sandbox = FileSystemSandbox(workspace_root=tmp_path)
    context = ToolExecutionContext(
        execution_id="test_xlsx_2",
        tool_id="office.create_xlsx",
        workspace_dir=tmp_path,
        sandbox=sandbox,
    )

    # Exceed column limit
    huge_headers = [f"Col{i}" for i in range(250)]
    sheets = [{"name": "WideSheet", "headers": huge_headers, "rows": []}]

    with pytest.raises(ResourceLimitExceededError):
        await create_xlsx(
            filename="too_wide.xlsx",
            sheets=sheets,
            context=context,
        )
