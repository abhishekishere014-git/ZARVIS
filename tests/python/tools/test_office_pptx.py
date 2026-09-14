"""Tests for native PPTX generation and presentation verification."""

from pathlib import Path
import pptx
import pytest
from jarvis.tools.builtins.office.pptx_tool import create_pptx
from jarvis.tools.errors import ResourceLimitExceededError
from jarvis.tools.models import ToolExecutionContext
from jarvis.tools.sandbox import FileSystemSandbox


@pytest.mark.asyncio
async def test_create_pptx_full_lifecycle(tmp_path: Path) -> None:
    sandbox = FileSystemSandbox(workspace_root=tmp_path)
    context = ToolExecutionContext(
        execution_id="test_pptx_1",
        tool_id="office.create_pptx",
        workspace_dir=tmp_path,
        sandbox=sandbox,
    )

    slides = [
        {
            "title": "Architecture Overview",
            "bullets": [
                "Python Core: Intelligence & Orchestration",
                "Node Gateway: Client real-time streaming",
                "Shared Protocol: Versioned contracts",
            ],
        },
        {
            "title": "Phase 04 Deliverables",
            "content": "Fully verified and sandboxed tool registry.",
            "bullets": [
                "Automatic schema introspection",
                "Native Office generation suite",
                "Secret scrubbed audit trail",
            ],
        },
    ]

    result = await create_pptx(
        filename="architecture_deck.pptx",
        title="JARVIS System Architecture",
        subtitle="Technical Deep Dive",
        slides=slides,
        context=context,
    )

    assert result["status"] == "success"
    assert result["filename"] == "architecture_deck.pptx"
    assert result["slides_count"] == 3  # 1 title + 2 content slides

    # Verify physical file and presentation structure
    file_path = Path(result["path"])
    assert file_path.exists()

    prs = pptx.Presentation(str(file_path))
    assert len(prs.slides) == 3
    assert prs.slides[0].shapes.title.text == "JARVIS System Architecture"
    assert prs.slides[1].shapes.title.text == "Architecture Overview"


@pytest.mark.asyncio
async def test_create_pptx_enforces_limits(tmp_path: Path) -> None:
    sandbox = FileSystemSandbox(workspace_root=tmp_path)
    context = ToolExecutionContext(
        execution_id="test_pptx_2",
        tool_id="office.create_pptx",
        workspace_dir=tmp_path,
        sandbox=sandbox,
    )

    # Exceed slide limit (> 100)
    huge_slides = [{"title": f"Slide {i}", "bullets": ["content"]} for i in range(105)]

    with pytest.raises(ResourceLimitExceededError):
        await create_pptx(
            filename="too_many_slides.pptx",
            title="Massive Deck",
            slides=huge_slides,
            context=context,
        )
